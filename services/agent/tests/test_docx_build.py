from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Mm
from jsonschema import ValidationError

from atlas_agent.core.events import ToolResult
from atlas_agent.core.policy import Tier, decide
from atlas_agent.llm.base import AgentTurn, ToolCall
from atlas_agent.tools.context import ToolContext, ToolError
from atlas_agent.tools.registry import ToolRegistry
from tests.conftest import (
    SequenceProvider,
    make_junction,
    outside_layout,
    run_scripted,
)

TIMES = "Times New Roman"
COURIER = "Courier New"


def _ctx(workspace: Path) -> ToolContext:
    return ToolContext(workspace=workspace, run_id="run-docx")


def _blob(payload: object) -> str:
    if isinstance(payload, BaseException):
        return str(payload)
    return json.dumps(payload, default=str, ensure_ascii=False)


def assert_no_workspace_path(workspace: Path | None, payload: object) -> None:
    blob = _blob(payload)
    if workspace is None:
        return
    resolved = str(workspace.resolve())
    assert resolved not in blob
    assert resolved.replace("\\", "/") not in blob
    assert resolved.replace("/", "\\") not in blob


def _meta(**overrides: object) -> dict[str, object]:
    data: dict[str, object] = {
        "student_name": "Test Student",
        "student_id": "00000",
        "instructor": "Test Instructor",
        "subject": "Intro to Testing",
        "deadline": "2099-01-01",
        "total_marks": "10",
        "assignment_number": 2,
    }
    data.update(overrides)
    return data


def _args(path: str = "out.docx", **overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "path": path,
        "meta": _meta(),
        "questions": [
            {
                "number": "Q1",
                "title": "Short answer",
                "marks": 1,
                "blocks": [
                    {"type": "paragraph", "text": "Read the prompt."},
                    {"type": "labeled", "label": "Answer:", "text": "A unit test."},
                    {
                        "type": "bullets",
                        "items": ["First point", "Second point"],
                    },
                    {
                        "type": "code",
                        "text": "def ping():\n    return 'pong'",
                    },
                ],
            },
            {
                "number": "Q2",
                "title": "Table question",
                "marks": 5,
                "blocks": [
                    {"type": "paragraph", "text": "Complete the table."},
                    {
                        "type": "table",
                        "headers": ["Case", "Result"],
                        "rows": [
                            ["ok", "pass"],
                            ["warn", "maybe"],
                            ["info", "note"],
                            ["bad", "fail"],
                        ],
                        "tones": [
                            ["good", "none"],
                            ["warn", "none"],
                            ["info", "none"],
                            ["bad", "none"],
                        ],
                    },
                ],
            },
        ],
    }
    payload.update(overrides)
    return payload


async def _run(registry: ToolRegistry, args: dict[str, object], ctx: ToolContext) -> dict[str, object]:
    return await registry.run("docx.build", args, ctx)


def _typeface(style: object) -> dict[str, str | None]:
    element = style.element  # type: ignore[attr-defined]
    r_pr = element.rPr
    assert r_pr is not None
    r_fonts = r_pr.rFonts
    assert r_fonts is not None
    return {
        "ascii": r_fonts.get(qn("w:ascii")),
        "hAnsi": r_fonts.get(qn("w:hAnsi")),
        "cs": r_fonts.get(qn("w:cs")),
        "eastAsia": r_fonts.get(qn("w:eastAsia")),
    }


def test_happy_path_header_footer_fonts_and_page_size(tmp_path: Path) -> None:
    workspace = tmp_path / "ws"
    workspace.mkdir()
    registry = ToolRegistry.load()
    ctx = _ctx(workspace)

    async def _call() -> None:
        letter = await _run(registry, _args("letter.docx"), ctx)
        assert letter["path"] == "letter.docx"
        assert letter["questions"] == 2
        assert letter["tables"] == 1
        assert letter["bytes"] > 0
        assert_no_workspace_path(workspace, letter)
        doc = Document(workspace / "letter.docx")
        header = doc.sections[0].header.paragraphs[0].text
        footer = doc.sections[0].footer.paragraphs[0].text
        assert header == "Intro to Testing — Assignment #2 | Iqra University Islamabad"
        assert footer == (
            "Test Student | Student ID: 00000 | Test Instructor | Iqra University Islamabad"
        )
        headings = [p.text for p in doc.paragraphs if p.style.name == "QuestionHeading"]
        assert headings[0] == "Q1. Short answer (1\u00a0Mark)"
        assert headings[1] == "Q2. Table question (5\u00a0Marks)"
        for name in ("Normal", "Body", "QuestionHeading", "SubLabel", "TableText"):
            faces = _typeface(doc.styles[name])
            assert set(faces.values()) == {TIMES}
        assert set(_typeface(doc.styles["Code"]).values()) == {COURIER}
        assert doc.sections[0].page_width == Inches(8.5)
        assert doc.sections[0].page_height == Inches(11)
        a4 = await _run(registry, _args("a4.docx", page_size="a4"), ctx)
        assert a4["path"] == "a4.docx"
        a4_doc = Document(workspace / "a4.docx")
        assert abs(int(a4_doc.sections[0].page_width) - int(Mm(210))) < 500
        assert abs(int(a4_doc.sections[0].page_height) - int(Mm(297))) < 500

    asyncio.run(_call())


def test_tables_repeat_header_cant_split_no_height_tones(tmp_path: Path) -> None:
    workspace = tmp_path / "ws"
    workspace.mkdir()
    registry = ToolRegistry.load()
    ctx = _ctx(workspace)

    async def _call() -> None:
        await _run(registry, _args("tables.docx"), ctx)
        doc = Document(workspace / "tables.docx")
        xml = doc.element.body.xml
        assert "trHeight" not in xml
        content = doc.tables[-1]
        header = content.rows[0]
        assert header._tr.trPr is not None
        assert header._tr.trPr.find(qn("w:tblHeader")) is not None
        for row in content.rows:
            assert row._tr.trPr is not None
            assert row._tr.trPr.find(qn("w:cantSplit")) is not None
            assert row._tr.trPr.find(qn("w:trHeight")) is None
        fills = []
        for row in content.rows[1:]:
            shd = row.cells[0]._tc.tcPr.find(qn("w:shd"))
            fills.append(shd.get(qn("w:fill")))
        assert fills == ["E9F7EF", "FEF9E7", "EBF5FB", "FDECEA"]
        bad = _args("bad-row.docx")
        questions = bad["questions"]
        assert isinstance(questions, list)
        q2 = questions[1]
        assert isinstance(q2, dict)
        q2["blocks"] = [
            {
                "type": "table",
                "headers": ["A", "B"],
                "rows": [["only-one"]],
            }
        ]
        with pytest.raises(ToolError, match="row length") as exc_info:
            await _run(registry, bad, ctx)
        assert_no_workspace_path(workspace, str(exc_info.value))

    asyncio.run(_call())


def test_code_shading_and_bullet_list_style(tmp_path: Path) -> None:
    workspace = tmp_path / "ws"
    workspace.mkdir()
    registry = ToolRegistry.load()
    ctx = _ctx(workspace)

    async def _call() -> None:
        await _run(registry, _args("blocks.docx"), ctx)
        doc = Document(workspace / "blocks.docx")
        code_paras = [p for p in doc.paragraphs if p.style.name == "Code"]
        assert code_paras
        shd = code_paras[0]._p.pPr.find(qn("w:shd"))
        assert shd is not None
        assert shd.get(qn("w:fill")) == "F4F4F4"
        assert set(_typeface(doc.styles["Code"]).values()) == {COURIER}
        bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
        assert [p.text for p in bullets] == ["First point", "Second point"]

    asyncio.run(_call())


def test_jail_rejects_unsafe_paths_without_absolute_errors(tmp_path: Path) -> None:
    workspace, outside = outside_layout(tmp_path)
    registry = ToolRegistry.load()
    ctx = _ctx(workspace)
    existing = workspace / "keep.docx"
    existing.write_bytes(b"PK-not-really")
    before = existing.read_bytes()

    async def _call() -> None:
        cases = [
            _args(str(workspace / "abs.docx")),
            _args("../escape.docx"),
            _args("C:/windows/out.docx"),
            _args("_trash/out.docx"),
            _args("notes.txt"),
            _args("keep.docx"),
        ]
        for args in cases:
            with pytest.raises(ToolError) as exc_info:
                await _run(registry, args, ctx)
            assert_no_workspace_path(workspace, str(exc_info.value))
            assert str(workspace.resolve()) not in str(exc_info.value)
        assert existing.read_bytes() == before
        link = workspace / "junc"
        make_junction(link, outside)
        with pytest.raises(ToolError) as junction:
            await _run(registry, _args("junc/out.docx"), ctx)
        assert_no_workspace_path(workspace, str(junction.value))
        assert "junc" in str(junction.value) or "links" in str(junction.value)
        try:
            (workspace / "sym").symlink_to(outside, target_is_directory=True)
        except OSError:
            pass
        else:
            with pytest.raises(ToolError) as symlink:
                await _run(registry, _args("sym/out.docx"), ctx)
            assert_no_workspace_path(workspace, str(symlink.value))

    asyncio.run(_call())


def test_limits_questions_strings_and_controls(tmp_path: Path) -> None:
    workspace = tmp_path / "ws"
    workspace.mkdir()
    registry = ToolRegistry.load()
    ctx = _ctx(workspace)

    async def _call() -> None:
        too_many = _args("many.docx")
        too_many["questions"] = [
            {
                "number": f"Q{i}",
                "title": "x",
                "marks": 1,
                "blocks": [{"type": "paragraph", "text": "n"}],
            }
            for i in range(31)
        ]
        with pytest.raises(ValidationError):
            await _run(registry, too_many, ctx)
        oversized = _args("big.docx")
        questions = oversized["questions"]
        assert isinstance(questions, list)
        q0 = questions[0]
        assert isinstance(q0, dict)
        q0["title"] = "a" * 2001
        with pytest.raises(ValidationError):
            await _run(registry, oversized, ctx)
        control = _args("ctrl.docx")
        qs = control["questions"]
        assert isinstance(qs, list)
        first = qs[0]
        assert isinstance(first, dict)
        first["blocks"] = [{"type": "paragraph", "text": "bad\x01char"}]
        with pytest.raises(ToolError, match="control") as exc_info:
            await _run(registry, control, ctx)
        assert_no_workspace_path(workspace, str(exc_info.value))

    asyncio.run(_call())


def test_registry_tier_policy_and_loop_validates_schema(tmp_path: Path) -> None:
    registry = ToolRegistry.load()
    loaded = registry.get("docx.build")
    assert loaded is not None
    assert loaded.spec.tier is Tier.T1
    decision = decide(loaded.spec, _args())
    assert decision.action == "allow"

    async def _call() -> None:
        workspace = tmp_path / "ws"
        workspace.mkdir()
        provider = SequenceProvider(
            [
                AgentTurn(
                    tool_calls=[
                        ToolCall(
                            name="docx.build",
                            args={"path": "skip.docx", "unexpected": True},
                        )
                    ]
                ),
                AgentTurn(text="done"),
            ]
        )
        emitter = await run_scripted(
            "build it",
            tmp_path,
            provider=provider,
            workspace=workspace,
        )
        result = emitter.of_type("tool_result")[0]
        assert isinstance(result, ToolResult)
        assert result.ok is False
        assert result.error is not None
        assert "invalid args" in result.error
        assert not (workspace / "skip.docx").exists()

    asyncio.run(_call())
