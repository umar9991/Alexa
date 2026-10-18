from __future__ import annotations

import asyncio
import logging
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Any, cast

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.table import Table, _Row
from docx.text.font import Font
from docx.text.paragraph import Paragraph

from atlas_agent.core.pathjail import display_path, resolve_in_workspace
from atlas_agent.tools.context import ToolContext, ToolError, require_workspace

logger = logging.getLogger("atlas_agent")

TIMES = "Times New Roman"
COURIER = "Courier New"
NBSP = "\u00a0"
MAX_BLOCKS = 200
MAX_COLS = 6
MAX_TABLE_ROWS = 30
MAX_STR = 2000
MAX_CODE = 4000
DEFAULT_UNIVERSITY = "IQRA UNIVERSITY ISLAMABAD"
DEFAULT_DEPARTMENT = "Department of Computer Science"
DEFAULT_UNI_SHORT = "Iqra University Islamabad"
FILL_BANNER = "2C3E50"
FILL_INFO = "F2F2F2"
FILL_HEADER = "F2F2F2"
FILL_CODE = "F4F4F4"
BORDER_GREY = "BFBFBF"
MUTED = RGBColor(0x54, 0x54, 0x54)
TONE_FILL = {
    "none": None,
    "good": "E9F7EF",
    "warn": "FEF9E7",
    "info": "EBF5FB",
    "bad": "FDECEA",
}


async def run(args: dict[str, object], ctx: ToolContext) -> dict[str, object]:
    workspace = require_workspace(ctx)
    return await asyncio.to_thread(_build_sync, workspace, args)


def _build_sync(workspace: Path, args: dict[str, object]) -> dict[str, object]:
    rel = args.get("path", "")
    if not isinstance(rel, str):
        raise ToolError("path must be a string")
    target = resolve_in_workspace(workspace, rel, allow_root=False, forbid_trash=True)
    shown = display_path(workspace, target)
    if not shown.lower().endswith(".docx"):
        raise ToolError("path must end with .docx")
    if target.exists():
        raise ToolError("file exists; overwriting is not available")
    meta = _meta(args.get("meta"))
    page_size = args.get("page_size", "letter")
    if page_size not in ("letter", "a4"):
        raise ToolError("page_size must be letter or a4")
    questions = _questions(args.get("questions"))
    try:
        payload = _render(meta, str(page_size), questions)
    except ToolError:
        raise
    except Exception:
        logger.exception("docx.build failed")
        raise ToolError("internal tool error") from None
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        with target.open("xb") as handle:
            handle.write(payload)
    except FileExistsError:
        raise ToolError("file exists; overwriting is not available") from None
    except OSError:
        logger.exception("docx.build write failed")
        raise ToolError("internal tool error") from None
    tables = sum(1 for q in questions for b in q["blocks"] if b["type"] == "table")
    return {
        "path": shown,
        "questions": len(questions),
        "tables": tables,
        "bytes": len(payload),
    }


def _meta(raw: object) -> dict[str, object]:
    if not isinstance(raw, dict):
        raise ToolError("meta is required")
    required = (
        "student_name",
        "student_id",
        "instructor",
        "subject",
        "deadline",
        "total_marks",
        "assignment_number",
    )
    out: dict[str, object] = {}
    for key in required:
        value = raw.get(key)
        if key == "assignment_number":
            if not isinstance(value, int) or isinstance(value, bool):
                raise ToolError("assignment_number must be an integer")
            out[key] = value
            continue
        if not isinstance(value, str):
            raise ToolError(f"{key} must be a string")
        _reject_controls(value, allow_newline=False, allow_tab=False)
        _limit_str(value, MAX_STR)
        out[key] = value
    for key, default in (
        ("university", DEFAULT_UNIVERSITY),
        ("department", DEFAULT_DEPARTMENT),
        ("university_short", DEFAULT_UNI_SHORT),
    ):
        value = raw.get(key, default)
        if value is None or value == "":
            value = default
        if not isinstance(value, str):
            raise ToolError(f"{key} must be a string")
        _reject_controls(value, allow_newline=False, allow_tab=False)
        _limit_str(value, MAX_STR)
        out[key] = value
    return out


def _questions(raw: object) -> list[dict[str, Any]]:
    if not isinstance(raw, list) or not raw:
        raise ToolError("questions must be a non-empty array")
    if len(raw) > 30:
        raise ToolError("too many questions")
    total_blocks = 0
    parsed: list[dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            raise ToolError("each question must be an object")
        number = item.get("number")
        title = item.get("title")
        marks = item.get("marks", None)
        blocks = item.get("blocks")
        if not isinstance(number, str) or not isinstance(title, str):
            raise ToolError("question number and title must be strings")
        _reject_controls(number, allow_newline=False, allow_tab=False)
        _reject_controls(title, allow_newline=False, allow_tab=False)
        _limit_str(number, MAX_STR)
        _limit_str(title, MAX_STR)
        if marks is not None and (not isinstance(marks, int) or isinstance(marks, bool)):
            raise ToolError("marks must be an integer or null")
        if not isinstance(blocks, list):
            raise ToolError("blocks must be an array")
        total_blocks += len(blocks)
        if total_blocks > MAX_BLOCKS:
            raise ToolError("too many blocks")
        parsed.append(
            {
                "number": number,
                "title": title,
                "marks": marks,
                "blocks": [_block(block) for block in blocks],
            }
        )
    return parsed


def _block(raw: object) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ToolError("each block must be an object")
    kind = raw.get("type")
    if kind == "paragraph":
        text = raw.get("text")
        if not isinstance(text, str):
            raise ToolError("paragraph text must be a string")
        _reject_controls(text, allow_newline=True, allow_tab=False)
        _limit_str(text, MAX_STR)
        return {"type": "paragraph", "text": text}
    if kind == "labeled":
        label = raw.get("label")
        text = raw.get("text")
        if not isinstance(label, str) or not isinstance(text, str):
            raise ToolError("labeled block needs label and text")
        _reject_controls(label, allow_newline=False, allow_tab=False)
        _reject_controls(text, allow_newline=True, allow_tab=False)
        _limit_str(label, MAX_STR)
        _limit_str(text, MAX_STR)
        return {"type": "labeled", "label": label, "text": text}
    if kind == "bullets":
        items = raw.get("items")
        if not isinstance(items, list):
            raise ToolError("bullets items must be an array")
        out: list[str] = []
        for item in items:
            if not isinstance(item, str):
                raise ToolError("bullet items must be strings")
            _reject_controls(item, allow_newline=False, allow_tab=False)
            _limit_str(item, MAX_STR)
            out.append(item)
        return {"type": "bullets", "items": out}
    if kind == "table":
        return _table_block(raw)
    if kind == "code":
        text = raw.get("text")
        if not isinstance(text, str):
            raise ToolError("code text must be a string")
        _reject_controls(text, allow_newline=True, allow_tab=True)
        if len(text) > MAX_CODE:
            raise ToolError("code is too long")
        return {"type": "code", "text": text}
    raise ToolError("unknown block type")


def _table_block(raw: dict[str, Any]) -> dict[str, Any]:
    headers = raw.get("headers")
    rows = raw.get("rows")
    tones = raw.get("tones")
    if not isinstance(headers, list) or not headers:
        raise ToolError("table headers are required")
    if len(headers) > MAX_COLS:
        raise ToolError("too many table columns")
    cols = len(headers)
    header_vals: list[str] = []
    for cell in headers:
        if not isinstance(cell, str):
            raise ToolError("table headers must be strings")
        _reject_controls(cell, allow_newline=False, allow_tab=False)
        _limit_str(cell, MAX_STR)
        header_vals.append(cell)
    if not isinstance(rows, list):
        raise ToolError("table rows must be an array")
    if len(rows) > MAX_TABLE_ROWS:
        raise ToolError("too many table rows")
    row_vals: list[list[str]] = []
    for row in rows:
        if not isinstance(row, list):
            raise ToolError("each table row must be an array")
        if len(row) != cols:
            raise ToolError("table row length must match headers")
        cells: list[str] = []
        for cell in row:
            if not isinstance(cell, str):
                raise ToolError("table cells must be strings")
            _reject_controls(cell, allow_newline=True, allow_tab=False)
            _limit_str(cell, MAX_STR)
            cells.append(cell)
        row_vals.append(cells)
    tone_vals: list[list[str]] | None = None
    if tones is not None:
        if not isinstance(tones, list) or len(tones) != len(row_vals):
            raise ToolError("table tones must match rows")
        tone_vals = []
        for row in tones:
            if not isinstance(row, list) or len(row) != cols:
                raise ToolError("table tones must match rows")
            painted: list[str] = []
            for cell in row:
                if cell not in TONE_FILL:
                    raise ToolError("invalid table tone")
                painted.append(str(cell))
            tone_vals.append(painted)
    return {
        "type": "table",
        "headers": header_vals,
        "rows": row_vals,
        "tones": tone_vals,
    }


def _limit_str(value: str, limit: int) -> None:
    if len(value) > limit:
        raise ToolError("string is too long")


def _reject_controls(value: str, *, allow_newline: bool, allow_tab: bool) -> None:
    for char in value:
        if char == "\n" and allow_newline:
            continue
        if char == "\t" and allow_tab:
            continue
        if unicodedata.category(char) == "Cc":
            raise ToolError("control characters are not allowed")


def _render(meta: dict[str, object], page_size: str, questions: list[dict[str, Any]]) -> bytes:
    doc = Document()
    _configure_styles(doc)
    _configure_section(doc, page_size)
    _configure_header_footer(doc, meta)
    _write_cover(doc, meta)
    for question in questions:
        _write_question(doc, question)
    end = doc.add_paragraph("— End of Assignment —", style="Body")
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in end.runs:
        run.italic = True
    subject = cast(str, meta["subject"])
    number = cast(int, meta["assignment_number"])
    doc.core_properties.title = f"{subject} — Assignment #{number}"
    doc.core_properties.author = cast(str, meta["student_name"])
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _configure_styles(doc: DocumentObject) -> None:
    styles = doc.styles
    normal = cast(ParagraphStyle, styles["Normal"])
    _apply_typeface(normal.font, TIMES)
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)

    body = _paragraph_style(doc, "Body")
    body.base_style = normal
    _apply_typeface(body.font, TIMES)
    body.font.size = Pt(11)
    body.paragraph_format.line_spacing = 1.1
    body.paragraph_format.space_after = Pt(6)

    heading2 = styles["Heading 2"]
    question = _paragraph_style(doc, "QuestionHeading")
    question.base_style = heading2
    _apply_typeface(question.font, TIMES)
    question.font.size = Pt(12)
    question.font.bold = True
    question.font.underline = True
    question.paragraph_format.space_before = Pt(14)
    question.paragraph_format.space_after = Pt(6)
    question.paragraph_format.keep_with_next = True
    _set_outline_level(question, 1)

    sub = _paragraph_style(doc, "SubLabel")
    sub.base_style = body
    _apply_typeface(sub.font, TIMES)
    sub.font.size = Pt(11)

    code = _paragraph_style(doc, "Code")
    code.base_style = body
    _apply_typeface(code.font, COURIER)
    code.font.size = Pt(10)
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)

    table_text = _paragraph_style(doc, "TableText")
    table_text.base_style = body
    _apply_typeface(table_text.font, TIMES)
    table_text.font.size = Pt(10.5)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.line_spacing = 1.1

    for name in ("Header", "Footer"):
        built = cast(ParagraphStyle, styles[name])
        _apply_typeface(built.font, TIMES)
        built.font.size = Pt(9)
        built.font.color.rgb = MUTED
        built.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _paragraph_style(doc: DocumentObject, name: str) -> ParagraphStyle:
    styles = doc.styles
    try:
        existing = styles[name]
        return cast(ParagraphStyle, existing)
    except KeyError:
        return cast(ParagraphStyle, styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH))


def _apply_typeface(font: Font, name: str) -> None:
    font.name = name
    rPr = font.element.get_or_add_rPr() if hasattr(font.element, "get_or_add_rPr") else font.element.rPr
    if rPr is None:
        return
    r_fonts = rPr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:cs"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def _set_outline_level(style: ParagraphStyle, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    tag = qn("w:outlineLvl")
    for old in p_pr.findall(tag):
        p_pr.remove(old)
    node = OxmlElement("w:outlineLvl")
    node.set(qn("w:val"), str(level))
    p_pr.append(node)


def _configure_section(doc: DocumentObject, page_size: str) -> None:
    section = doc.sections[0]
    if page_size == "a4":
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.different_first_page_header_footer = False


def _configure_header_footer(doc: DocumentObject, meta: dict[str, object]) -> None:
    section = doc.sections[0]
    subject = cast(str, meta["subject"])
    number = cast(int, meta["assignment_number"])
    short = cast(str, meta["university_short"])
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.style = doc.styles["Header"]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.text = f"{subject} — Assignment #{number} | {short}"
    _set_paragraph_border(hp, edge="bottom")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.style = doc.styles["Footer"]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    student = cast(str, meta["student_name"])
    student_id = cast(str, meta["student_id"])
    instructor = cast(str, meta["instructor"])
    fp.text = f"{student} | Student ID: {student_id} | {instructor} | {short}"
    _set_paragraph_border(fp, edge="top")


def _set_paragraph_border(paragraph: Paragraph, *, edge: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    tag = qn("w:pBdr")
    for old in p_pr.findall(tag):
        p_pr.remove(old)
    borders = OxmlElement("w:pBdr")
    line = OxmlElement(f"w:{edge}")
    line.set(qn("w:val"), "single")
    line.set(qn("w:sz"), "6")
    line.set(qn("w:space"), "1")
    line.set(qn("w:color"), BORDER_GREY)
    borders.append(line)
    p_pr.append(borders)


def _write_cover(doc: DocumentObject, meta: dict[str, object]) -> None:
    university = doc.add_paragraph(style="Body")
    university.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = university.add_run(cast(str, meta["university"]))
    run.bold = True
    run.font.size = Pt(14)

    dept = doc.add_paragraph(cast(str, meta["department"]), style="Body")
    dept.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subject = cast(str, meta["subject"]).upper()
    number = cast(int, meta["assignment_number"])
    banner = doc.add_table(rows=1, cols=1)
    _set_table_borders(banner, BORDER_GREY)
    _set_table_full_width(doc, banner)
    cell = banner.cell(0, 0)
    _shade_tc(cell._tc, FILL_BANNER)
    para = cell.paragraphs[0]
    para.style = doc.styles["TableText"]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text = para.add_run(f"{subject} — ASSIGNMENT #{number}")
    text.bold = True
    text.font.size = Pt(13)
    text.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    spacer = doc.add_paragraph("", style="Body")
    spacer.paragraph_format.space_after = Pt(4)

    info = doc.add_table(rows=3, cols=2)
    _set_table_borders(info, BORDER_GREY)
    _set_table_full_width(doc, info)
    pairs = (
        ("Student Name", cast(str, meta["student_name"]), "Student ID", cast(str, meta["student_id"])),
        ("Instructor", cast(str, meta["instructor"]), "Total Marks", cast(str, meta["total_marks"])),
        ("Subject", cast(str, meta["subject"]), "Deadline", cast(str, meta["deadline"])),
    )
    for row_idx, (l1, v1, l2, v2) in enumerate(pairs):
        row = info.rows[row_idx]
        _row_cant_split(row)
        _info_cell(row.cells[0], l1, v1)
        _info_cell(row.cells[1], l2, v2)


def _info_cell(cell: Any, label: str, value: str) -> None:
    _shade_tc(cell._tc, FILL_INFO)
    first = cell.paragraphs[0]
    first.style = "TableText"
    label_run = first.add_run(label)
    label_run.bold = True
    value_p = cell.add_paragraph(value, style="TableText")
    for paragraph in (first, value_p):
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(2)


def _write_question(doc: DocumentObject, question: dict[str, Any]) -> None:
    heading = _question_heading_text(question)
    para = doc.add_paragraph(heading, style="QuestionHeading")
    para.paragraph_format.keep_with_next = True
    last_para = para
    for block in question["blocks"]:
        kind = block["type"]
        if kind == "paragraph":
            last_para = doc.add_paragraph(block["text"], style="Body")
        elif kind == "labeled":
            last_para = doc.add_paragraph(style="SubLabel")
            label_run = last_para.add_run(block["label"])
            label_run.bold = True
            last_para.add_run(" " + block["text"])
        elif kind == "bullets":
            for item in block["items"]:
                last_para = doc.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            last_para.paragraph_format.keep_with_next = True
            last_para = _write_table(doc, block)
        elif kind == "code":
            last_para = _write_code(doc, block["text"])


def _question_heading_text(question: dict[str, Any]) -> str:
    base = f"{question['number']}. {question['title']}"
    marks = question["marks"]
    if marks is None:
        return base
    word = "Mark" if marks == 1 else "Marks"
    return f"{base} ({marks}{NBSP}{word})"


def _write_code(doc: DocumentObject, text: str) -> Paragraph:
    para = doc.add_paragraph(style="Code")
    _shade_p(para, FILL_CODE)
    lines = text.split("\n")
    for index, line in enumerate(lines):
        run = para.add_run(line)
        if index < len(lines) - 1:
            run.add_break()
    return para


def _write_table(doc: DocumentObject, block: dict[str, Any]) -> Paragraph:
    headers: list[str] = block["headers"]
    rows: list[list[str]] = block["rows"]
    tones: list[list[str]] | None = block["tones"]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _set_table_borders(table, BORDER_GREY)
    _set_table_full_width(doc, table)
    header_row = table.rows[0]
    _row_cant_split(header_row)
    _row_repeat_header(header_row)
    for idx, title in enumerate(headers):
        cell = header_row.cells[idx]
        _shade_tc(cell._tc, FILL_HEADER)
        para = cell.paragraphs[0]
        para.style = "TableText"
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        run = para.add_run(title)
        run.bold = True
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        _row_cant_split(row)
        for c_idx, value in enumerate(row_data):
            cell = row.cells[c_idx]
            fill = None
            if tones is not None:
                fill = TONE_FILL[tones[r_idx][c_idx]]
            if fill:
                _shade_tc(cell._tc, fill)
            para = cell.paragraphs[0]
            para.style = "TableText"
            para.text = value
    return doc.add_paragraph("", style="Body")


def _set_table_full_width(doc: DocumentObject, table: Table) -> None:
    section = doc.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin
    table.autofit = True
    table.width = usable
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tag = qn("w:tblW")
    for old in tbl_pr.findall(tag):
        tbl_pr.remove(old)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    tbl_pr.append(width)


def _set_table_borders(table: Table, color: str) -> None:
    tbl_pr = table._tbl.tblPr
    tag = qn("w:tblBorders")
    for old in tbl_pr.findall(tag):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:space"), "0")
        line.set(qn("w:color"), color)
        borders.append(line)
    tbl_pr.append(borders)


def _row_cant_split(row: _Row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tag = qn("w:cantSplit")
    for old in tr_pr.findall(tag):
        tr_pr.remove(old)
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def _row_repeat_header(row: _Row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tag = qn("w:tblHeader")
    for old in tr_pr.findall(tag):
        tr_pr.remove(old)
    node = OxmlElement("w:tblHeader")
    tr_pr.append(node)


def _shade_tc(tc: Any, fill: str) -> None:
    tc_pr = tc.get_or_add_tcPr()
    _set_shd(tc_pr, fill)


def _shade_p(paragraph: Paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    _set_shd(p_pr, fill)


def _set_shd(parent: Any, fill: str) -> None:
    tag = qn("w:shd")
    for old in parent.findall(tag):
        parent.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    parent.append(shd)
